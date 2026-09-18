"""Logic xử lý chuyển đổi bảng từ DOCX sang Excel."""
import io
import os
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

def inspect_docx_tables(file_source) -> list:
    """Đọc file DOCX (đường dẫn hoặc BytesIO) và trả về danh sách thông tin các bảng."""
    doc = Document(file_source)
    tables_info = []
    for idx, tbl in enumerate(doc.tables):
        rows_cnt = len(tbl.rows)
        cols_cnt = len(tbl.columns) if rows_cnt > 0 else 0
        first_row_preview = []
        if rows_cnt > 0:
            first_row_preview = [cell.text.strip().replace("\n", " ") for cell in tbl.rows[0].cells[:4]]
            if len(tbl.rows[0].cells) > 4:
                first_row_preview.append("...")
        tables_info.append({
            "index": idx,
            "rows": rows_cnt,
            "cols": cols_cnt,
            "preview": " | ".join(first_row_preview) if first_row_preview else "(Bảng rỗng)"
        })
    return tables_info

def convert_docx_table_to_excel_buffer(file_source, table_index: int = 0) -> io.BytesIO:
    """Chuyển đổi bảng từ DOCX sang workbook Excel dưới dạng in-memory BytesIO."""
    doc = Document(file_source)
    if not doc.tables:
        raise ValueError("File DOCX không chứa bảng nào.")

    if table_index >= len(doc.tables):
        raise IndexError(f"Chỉ định table_index={table_index} nhưng file chỉ có {len(doc.tables)} bảng.")

    target_table = doc.tables[table_index]
    wb = Workbook()
    ws = wb.active
    ws.title = f"Table_{table_index + 1}"

    thin_border_side = Side(style="thin", color="000000")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    header_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )
    header_font = Font(name="Arial", size=10, bold=True)
    body_font = Font(name="Arial", size=10)

    # Đọc và ghi dữ liệu từng ô
    for row_idx, row in enumerate(target_table.rows, start=1):
        for col_idx, cell in enumerate(row.cells, start=1):
            excel_cell = ws.cell(row=row_idx, column=col_idx)

            cell_paragraphs = [p.text.strip() for p in cell.paragraphs]
            combined_text = "\n".join([text for text in cell_paragraphs if text])

            excel_cell.value = combined_text
            excel_cell.border = cell_border

            if row_idx == 1:
                excel_cell.font = header_font
                excel_cell.fill = header_fill
                excel_cell.alignment = Alignment(
                    wrap_text=True, vertical="center", horizontal="center"
                )
            else:
                excel_cell.font = body_font
                excel_cell.alignment = Alignment(
                    wrap_text=True, vertical="center", horizontal="left"
                )

    # Tự động tính toán độ rộng cột
    for col in ws.columns:
        max_line_len = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            if cell.value:
                lines = str(cell.value).split("\n")
                longest_line = max(len(line) for line in lines) if lines else 0
                max_line_len = max(max_line_len, longest_line)

        adjusted_width = max(max_line_len + 3, 14)
        adjusted_width = min(adjusted_width, 60)
        ws.column_dimensions[col_letter].width = adjusted_width

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream

def save_excel_to_local_path(excel_bytes: bytes, target_directory: str, filename: str) -> str:
    """Lưu dữ liệu Excel trực tiếp vào một thư mục trên máy tính/server."""
    clean_dir = os.path.expanduser(target_directory.strip().strip('"').strip("'"))
    if not os.path.exists(clean_dir):
        os.makedirs(clean_dir, exist_ok=True)
    
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"
        
    full_path = os.path.join(clean_dir, filename)
    with open(full_path, "wb") as f:
        f.write(excel_bytes)
    return full_path