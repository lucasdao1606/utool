import io
import openpyxl
from typing import List, Tuple
from .models import SpecTargetItem, ItemAuditResult

def parse_template_excel(file_bytes) -> Tuple[openpyxl.Workbook, List[SpecTargetItem]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active

    items = []
    for r in range(2, ws.max_row + 1):
        val_tt = ws.cell(r, 1).value
        if val_tt is None:
            continue

        try:
            tt_int = int(str(val_tt).strip())
        except (ValueError, TypeError):
            continue

        name = str(ws.cell(r, 2).value or "").strip()
        req_text = str(ws.cell(r, 3).value or "").strip()
        ctcb = str(ws.cell(r, 4).value or "").strip()
        manufacturer = str(ws.cell(r, 5).value or "").strip()
        part_number = str(ws.cell(r, 6).value or "").strip()

        sub_specs = []
        if req_text:
            for line in req_text.split("\n"):
                line_clean = line.strip().lstrip("-*+ \u2022").strip()
                if line_clean:
                    sub_specs.append(line_clean)

        items.append(SpecTargetItem(
            row_idx=r,
            tt=tt_int,
            name=name,
            req_text=req_text,
            sub_specs=sub_specs,
            ctcb=ctcb,
            manufacturer=manufacturer,
            part_number=part_number
        ))

    return wb, items

def save_audit_results_to_workbook(wb: openpyxl.Workbook, results: List[ItemAuditResult], is_mode_audit: bool = True) -> bytes:
    ws = wb.active

    if is_mode_audit:
        # --- CH? Ð? 1: ÐÁNH GIÁ CH? TIÊU K? THU?T ---
        # "Thông s? k? thu?t (Th?c t?)"
        ws.cell(1, 7).value = "Th\u00f4ng s\u1ed1 k\u1ef9 thu\u1eadt (Th\u1ef1c t\u1ebf)"
        # "Nh?n xét c?a chuyên gia"
        ws.cell(1, 8).value = "Nh\u1eadn x\u00e9t c\u1ee7a chuy\u00ean gia"
        # "Tham chi?u"
        ws.cell(1, 9).value = "Tham chi\u1ebfu"
        # "Ghi chú"
        ws.cell(1, 10).value = "Ghi ch\u00fa"
        # "Ð? xu?t hi?u ch?nh"
        ws.cell(1, 11).value = "\u0110\u1ec1 xu\u1ea5t hi\u1ec7u ch\u1ec9nh"

        for res in results:
            r = res.row_idx
            ws.cell(r, 7).value = res.thong_so_ky_thuat
            ws.cell(r, 8).value = res.nhan_xet
            ws.cell(r, 9).value = res.tham_chieu
            ws.cell(r, 10).value = res.ghi_chu
            ws.cell(r, 11).value = res.de_xuat
    else:
        # --- CH? Ð? 2: T? Ð?NG XÂY D?NG CH? TIÊU ---
        # "Yêu c?u k? thu?t (AI T? d?ng l?p)"
        ws.cell(1, 3).value = "Y\u00eau c\u1ea7u k\u1ef9 thu\u1eadt (AI T\u1ef1 \u0111\u1ed9ng l\u1eadp)"
        # "Tr?ng thái h? th?ng"
        ws.cell(1, 10).value = "Tr\u1ea1ng th\u00e1i h\u1ec7 th\u1ed1ng"

        for res in results:
            r = res.row_idx
            if res.status == "PASS":
                ws.cell(r, 3).value = res.de_xuat
                # "AI dã di?n thành công"
                ws.cell(r, 10).value = "AI \u0111\u00e3 \u0111i\u1ec1n th\u00e0nh c\u00f4ng"
            else:
                ws.cell(r, 10).value = res.ghi_chu

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()