import io
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Bảng quy đổi đơn vị chuẩn
UNIT_MAP = {
    "kg": ("mass", 1), "g": ("mass", 1e-3), "mg": ("mass", 1e-6),
    "m": ("length", 1), "dm": ("length", 1e-1), "cm": ("length", 1e-2), "mm": ("length", 1e-3),
    "v": ("voltage", 1), "mv": ("voltage", 1e-3), "kv": ("voltage", 1e3),
    "hz": ("freq", 1), "khz": ("freq", 1e3), "mhz": ("freq", 1e6), "ghz": ("freq", 1e9),
}

def normalize_text(s):
    return str(s).lower().replace("-", "").replace("_", "").replace(" ", "")

def parse_value_with_unit(value):
    if pd.isna(value):
        return None

    s = str(value).lower().strip().replace(",", ".")
    num = re.findall(r"[-+]?\d*\.?\d+", s)
    unit = re.findall(r"[a-zA-Z]+", s)

    if not num:
        return None

    val = float(num[0])
    unit = unit[-1] if unit else ""

    if unit in UNIT_MAP:
        _, factor = UNIT_MAP[unit]
        return val * factor

    return val

def parse_requirement(req):
    if pd.isna(req):
        return None

    s = str(req).lower().strip().replace(",", ".")

    def val(x):
        return parse_value_with_unit(x)

    if "<=" in s or "≤" in s:
        return ("max", val(s))
    if ">=" in s or "≥" in s:
        return ("min", val(s))

    if "~" in s:
        nums = re.findall(r"[-+]?\d*\.?\d+", s)
        if len(nums) >= 2:
            return ("range", float(nums[0]), float(nums[1]))

    if "±" in s:
        nums = re.findall(r"[-+]?\d*\.?\d+", s)
        if len(nums) >= 2:
            base = float(nums[0])
            tol = float(nums[1]) / 100
            return ("range", base * (1 - tol), base * (1 + tol))

    if re.search(r"[\\\[\]\+\*\?\d]", s):
        return ("regex", s)

    return ("text", s)

def check_criteria(req_parsed, value):
    if req_parsed is None:
        return "WARNING", "No requirement"

    t = req_parsed[0]

    # Khớp dạng chuỗi văn bản
    if t == "text":
        if pd.isna(value):
            return "FAIL", "Empty"
        val_norm = normalize_text(value)
        options = re.split(r"/|,|or", req_parsed[1])
        for opt in options:
            if normalize_text(opt) in val_norm:
                return "PASS", f"match {opt}"
        return "FAIL", "no match"

    # Khớp Regex
    if t == "regex":
        if re.search(req_parsed[1], str(value), re.IGNORECASE):
            return "PASS", "regex ok"
        return "FAIL", "regex fail"

    # Giá trị số kèm đơn vị
    val = parse_value_with_unit(value)
    if val is None:
        return "FAIL", "NaN"

    if t == "min":
        return ("PASS", "OK") if val >= req_parsed[1] else ("FAIL", "< min")
    if t == "max":
        return ("PASS", "OK") if val <= req_parsed[1] else ("FAIL", "> max")
    if t == "range":
        return ("PASS", "OK") if req_parsed[1] <= val <= req_parsed[2] else ("FAIL", "out")
    if t == "exact":
        return ("PASS", "OK") if val == req_parsed[1] else ("FAIL", "neq")

    return "WARNING", "unknown"

def process_checklist(file_bytes: bytes, col_name=1, col_req=2, sp_cols=None):
    """
    Xử lý file Excel checklist đối soát kỹ thuật.
    sp_cols: Dict chứa tên sản phẩm và index cột tương ứng, ví dụ {"SP1": 4, "SP2": 6, "SP3": 8}
    """
    if sp_cols is None:
        sp_cols = {"SP1": 4, "SP2": 6, "SP3": 8}

    df = pd.read_excel(io.BytesIO(file_bytes), header=None).dropna(how="all")

    results = []
    score = {sp: 0 for sp in sp_cols}
    total = 0

    for i in range(len(df)):
        name = df.iloc[i, col_name]
        req = df.iloc[i, col_req]

        if pd.isna(name):
            continue

        total += 1
        parsed = parse_requirement(req)
        row = {"Tiêu chí": name, "Yêu cầu": req}

        for label, col_idx in sp_cols.items():
            v = df.iloc[i, col_idx] if col_idx < df.shape[1] else None
            s, r = check_criteria(parsed, v)
            row[label] = v
            row[f"{label}_Kết quả"] = s

            if s == "PASS":
                score[label] += 1

        results.append(row)

    score_percent = {k: round(v / total * 100, 2) if total else 0 for k, v in score.items()}
    res_df = pd.DataFrame(results)

    # Tô màu kết quả và lưu vào buffer nhị phân
    output = io.BytesIO()
    res_df.to_excel(output, index=False)
    output.seek(0)

    wb = load_workbook(output)
    ws = wb.active

    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # Xác định các cột chứa chữ "Kết quả" để tô màu
    header = [cell.value for cell in ws[1]]
    result_col_indices = [idx + 1 for idx, val in enumerate(header) if val and "_Kết quả" in str(val)]

    for row in range(2, ws.max_row + 1):
        for col_idx in result_col_indices:
            cell = ws.cell(row=row, column=col_idx)
            if cell.value == "PASS":
                cell.fill = green
            elif cell.value == "FAIL":
                cell.fill = red

    final_output = io.BytesIO()
    wb.save(final_output)
    wb.close()

    return res_df, score_percent, final_output.getvalue()